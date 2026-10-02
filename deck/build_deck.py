"""Builds Formation_Damage_Drilling.pptx  (23 slides + 5 hidden backups)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from deckkit import *
from pptx.enum.dml import MSO_LINE_DASH_STYLE

FIG = "figures/"
EMU_IN = 914400


def pic(slide, name, x, y, w):
    """Insert a figure at a given width; returns (x, y, w, h) in inches."""
    from PIL import Image
    iw, ih = Image.open(FIG + name).size
    h = w * ih / iw
    slide.shapes.add_picture(FIG + name, Inches(x), Inches(y), Inches(w), Inches(h))
    return x, y, w, h


def fit_title_size(s, w=CONTENT_W, p=30.0, max_lines=2, minimum=24.0):
    size = p
    while size > minimum and wrapped_lines(s, w, size) > max_lines:
        size -= 0.5
    return size


def heading(slide, markup, y=0.34, p=32.0, w=None):
    """Slide heading; auto-shrinks to fit two lines inside the given width."""
    w = w or CONTENT_W
    size = p
    while size > 24.0 and wrapped_lines(strip_markup(markup), w, size) > 2:
        size -= 0.5
    box, tf = textbox(slide, M, y, w, 1.0)
    rich(para(tf, first=True, line=1.02), markup, size, NAVY, True)
    return size


def five_boxes(slide, items, y, h=2.2, x0=M, gap=0.28, with_rail=False,
               rail_y=None):
    w = (CONTENT_W - 4 * gap) / 5
    if with_rail:
        line(slide, x0 + w / 2, rail_y, x0 + 4 * (w + gap) + w / 2, rail_y, LGREY, 1.5)
    for i, (name, caption) in enumerate(items):
        x = x0 + i * (w + gap)
        if with_rail:
            rect(slide, x + w / 2 - 0.135, rail_y - 0.135, 0.27, 0.27,
                 fill=TEAL, line=WHITE, lw=1.5, shape=MSO_SHAPE.OVAL)
        rect(slide, x, y, w, h, fill=WHITE, line=BORDER, lw=1.25)
        text(slide, name, x + 0.12, y + 0.42, w - 0.24, 0.5, size=17, color=NAVY,
             bold=True, align=PP_ALIGN.CENTER, line=1.05)
        text(slide, caption, x + 0.12, y + 1.02, w - 0.24, 1.0, size=12.5,
             color=GREY, align=PP_ALIGN.CENTER, line=1.25)
    return w


def prs_new():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ============================================================ 1  TITLE
def s01(prs):
    s = blank(prs)
    rect(s, 0, 0, 0.22, SLIDE_H, fill=NAVY)
    rect(s, 0.90, 1.02, 1.55, 0.09, fill=TEAL)
    box, tf = textbox(s, 0.90, 1.48, 7.60, 2.60)
    rich(para(tf, first=True, line=1.03), "Formation Damage", 35, NAVY, True)
    rich(para(tf, line=1.03), "During Drilling:", 35, NAVY, True)
    rich(para(tf, line=1.03), "Mechanisms, Skin Quantification,", 26, TEALD, True)
    rich(para(tf, line=1.03), "and Mitigation Strategies", 26, TEALD, True)
    box, tf = textbox(s, 0.90, 4.35, 7.4, 1.9)
    rich(para(tf, first=True, after=9), "[Your Name]", 19, INK, True)
    rich(para(tf, after=3), "Supervisor: [Prof. Name]", 14, GREY)
    rich(para(tf, after=3), "Department of Petroleum and Gas Engineering · [University]", 14, GREY)
    rich(para(tf, after=3), "Seminar · [Month Year]", 14, GREY)
    # decorative wellbore motif
    for dia, col, lw in [(4.6, RGBColor(0xE6, 0xEF, 0xF4), 1.5),
                         (3.2, RGBColor(0xCF, 0xE4, 0xE7), 1.5),
                         (2.0, RGBColor(0xC2, 0xE4, 0xE6), 1.5),
                         (0.95, TEAL, 2.0)]:
        rect(s, 11.05 - dia / 2, 3.60 - dia / 2, dia, dia, fill=None, line=col,
             lw=lw, shape=MSO_SHAPE.OVAL)
    rect(s, 11.05 - 0.16, 3.60 - 0.16, 0.32, 0.32, fill=NAVY, line=None,
         shape=MSO_SHAPE.OVAL)
    text(s, "Image: [source / licence]", 0.90, 6.86, 6.0, 0.28, size=10, color=GREY,
         italic=True)
    text(s, "23 slides + 5 hidden backups", SLIDE_W - M - 3.0, 6.86, 3.0, 0.28,
         size=10, color=LGREY, align=PP_ALIGN.RIGHT)
    notes(s, "Welcome. This seminar covers how formation damage forms during drilling, "
             "how we quantify it with skin, and what we can do about it. "
             "Fill in your name, supervisor, department and date before presenting.")
    return s


# ============================================================ 2  IPR
def s02(prs):
    s = blank(prs)
    heading(s, "Same drawdown, but damage cuts the rate by more than half")
    text(s, "−55%", M, 1.50, 3.6, 1.1, size=60, color=RED, bold=True)
    text(s, "rate at the same drawdown", M, 2.72, 3.9, 0.35, size=15, color=GREY)
    line(s, M, 3.22, 4.05, 3.22, BORDER, 1.25)
    text(s, "4,270 → 1,910", M, 3.45, 3.9, 0.5, size=27, color=NAVY, bold=True)
    text(s, "STB/d, undamaged → damaged", M, 4.00, 3.9, 0.4, size=13.5, color=GREY)
    text(s, "(S = 0 → S = 8.55)", M, 4.34, 3.9, 0.4, size=13.5, color=GREY)
    text(s, "At p~wf~ = 2,500 psi, same drawdown.", M, 4.82, 4.0, 0.4, size=13,
         color=GREY)
    pic(s, "s2_ipr.png", 4.70, 1.62, 8.05)
    source(s, "Illustrative data (k = 100 md, h = 50 ft). Author's calculation.")
    pagenum(s, 2)
    notes(s, "Open with the punchline: same reservoir, same drawdown, only skin differs — "
             "and the damaged well delivers less than half the rate. "
             "Everything that follows explains where skin 8.5 comes from and what to do about it.")
    return s


# ============================================================ 3  PROBLEM + OBJECTIVES
def s03(prs):
    s = blank(prs)
    heading(s, "Damage is invisible while drilling but costly during production")
    rect(s, M, 1.85, CONTENT_W, 0.72, fill=PANEL2)
    rect(s, M, 1.85, 0.10, 0.72, fill=AMBER)
    text(s, "Problem: drilling fluids impair near-wellbore permeability", M + 0.30, 2.03,
         CONTENT_W - 0.5, 0.4, size=16, color=NAVY, bold=True)
    items = [("1. Mechanisms", "Explain how damage forms"),
             ("2. Quantification", "Measure the loss with skin"),
             ("3. Strategies", "Compare prevention and remediation")]
    gap = 0.235
    w = (CONTENT_W - 2 * gap) / 3
    for i, (head, body) in enumerate(items):
        x = M + i * (w + gap)
        rect(s, x, 3.05, w, 3.15, fill=WHITE, line=BORDER, lw=1.25)
        rect(s, x + w / 2 - 0.42, 3.41, 0.84, 0.84, fill=TEALP, line=TEAL, lw=1.5,
             shape=MSO_SHAPE.OVAL)
        text(s, str(i + 1), x + w / 2 - 0.42, 3.59, 0.84, 0.5, size=24, color=TEALD,
             bold=True, align=PP_ALIGN.CENTER)
        text(s, head, x + 0.2, 4.55, w - 0.4, 0.4, size=19, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
        text(s, body, x + 0.25, 5.08, w - 0.5, 0.8, size=15, color=GREY,
             align=PP_ALIGN.CENTER, line=1.25)
    pagenum(s, 3)
    notes(s, "Three questions frame the talk: how does damage form, how do we measure the "
             "loss, and what can we do about it? Stress that damage is invisible while "
             "drilling but shows up as lost rate once the well is on production.")
    return s


# ============================================================ 4  ROADMAP
def s04(prs):
    s = blank(prs)
    heading(s, "Five steps, from flow physics to my project plan")
    five_boxes(s, [("Fundamentals", "Flow near the well"),
                   ("Mechanisms", "How damage forms"),
                   ("Quantification", "Skin and productivity"),
                   ("Solutions", "Prevention and cure"),
                   ("Next steps", "Economics and project")],
               y=3.05, h=2.15, with_rail=True, rail_y=2.58)
    for i in range(5):
        w = (CONTENT_W - 4 * 0.28) / 5
        x = M + i * (w + 0.28) + w / 2
        text(s, str(i + 1), x - 0.135, 2.445, 0.27, 0.3, size=11, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER)
    pagenum(s, 4)
    notes(s, "Roadmap: near-well flow fundamentals, the mechanisms that create damage, "
             "quantification through skin, prevention and remediation, then economics and "
             "my graduation project plan.")
    return s


# ============================================================ 5  NEAR-WELLBORE FLOW
def s05(prs):
    s = blank(prs)
    heading(s, "The near-wellbore zone controls inflow")
    tracker(s, 1)
    bullets_from(s, ["Flow area shrinks toward the well",
                     "Pressure gradient is steepest near the wellbore",
                     "About 30% of drawdown occurs within 3 ft",
                     "Small damaged zone, large effect"],
                M, 2.00, 5.15, 4.40, size=17)
    pic(s, "s5_radial.png", 6.35, 2.02, 5.45)
    source(s, "Schematic by author. Undamaged pseudo-steady-state flow.")
    pagenum(s, 5)
    notes(s, "Flow converges into the well, so the pressure gradient is steepest right at the "
             "wellbore. About 30% of the drawdown from the drainage boundary is spent within "
             "3 ft of the well — a tiny volume that controls inflow.")
    return s


# ============================================================ 6  SKIN DEFINITION
def s06(prs):
    s = blank(prs)
    heading(s, "Skin factor measures deviation from ideal radial flow")
    tracker(s, 1)
    bullets_from(s, ["**S > 0**: damaged",
                     "**S = 0**: ideal, undisturbed",
                     "**S < 0**: stimulated (acid, fracture)",
                     "Skin converts disturbance into extra pressure drop"],
                M, 2.00, 4.65, 4.40, size=17)
    pic(s, "s6_profile.png", 5.40, 1.98, 7.15)
    rect(s, 5.40, 6.24, 7.15, 0.58, fill=PANEL2)
    text(s, "Δp~skin~ = 141.2 q B μ S / (k h)", 5.40, 6.37, 7.15, 0.4, size=17,
         color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    source(s, "Schematic by author. After van Everdingen (1953); Hawkins (1956).", y=6.95)
    pagenum(s, 6)
    notes(s, "Skin is a dimensionless pressure-drop knob: positive means damage, zero is "
             "ideal, negative means stimulation. It converts everything happening near the "
             "wellbore into one extra pressure drop, which is why it is so useful.")
    return s


# ============================================================ 7  RADIAL FLOW EQUATION
def s07(prs):
    s = blank(prs)
    heading(s, "In the radial-flow equation, skin adds directly to resistance")
    tracker(s, 1)
    rect(s, M, 2.05, CONTENT_W, 1.32, fill=PANEL2, line=BORDER, lw=1.0)
    text(s, "q  =  0.00708 k h (p̄ − p~wf~) / [ μ B ( ln(r~e~/r~w~) − 0.75 + [red]**S** ) ]",
         M + 0.2, 2.50, CONTENT_W - 0.4, 0.6, size=25, color=NAVY, align=PP_ALIGN.CENTER)
    legend = [("q", "oil rate (STB/d)"), ("k", "permeability (md)"),
              ("h", "net thickness (ft)"), ("p̄", "average reservoir pressure (psi)"),
              ("p~wf~", "flowing bottomhole pressure (psi)"), ("μ", "viscosity (cP)"),
              ("B", "formation volume factor"), ("r~e~, r~w~", "drainage / well radius (ft)"),
              ("S", "skin (dimensionless)")]
    cw = 3.12
    for i, (sym, desc) in enumerate(legend):
        r, c = divmod(i, 2)
        x = M + c * cw
        y = 3.62 + r * 0.62
        box, tf = textbox(s, x, y, cw - 0.18, 0.5)
        p = para(tf, first=True, line=1.1)
        rich(p, sym, 15, NAVY, True)
        rich(p, "   " + desc, 12, GREY)
    line(s, 7.02, 3.55, 7.02, 6.55, BORDER, 1.25)
    bullets_from(s, ["Skin adds to the geometric term",
                     "Here ln(r~e~/r~w~) − 0.75 ≈ 6.9",
                     "So S ≈ 7 halves the rate"],
                7.38, 3.55, 5.40, 3.0, size=17)
    source(s, "Pseudo-steady-state radial flow, field units. Economides et al.", y=6.95)
    pagenum(s, 7)
    notes(s, "Walk the equation. In the denominator, skin adds to the geometry term "
             "ln(re/rw) − 0.75, which is about 6.9 here — so a skin of roughly 7 doubles the "
             "resistance and halves the rate. Note that S is dimensionless.")
    return s


# ============================================================ 8  DAMAGE AT EVERY STAGE
def s08(prs):
    s = blank(prs)
    heading(s, "Damage can occur at every stage of the well's life")
    tracker(s, 1)
    five_boxes(s, [("Drilling", "Mud invasion"),
                   ("Cementing", "Cement filtrate"),
                   ("Completion", "Perforation damage,\ncompletion fluid"),
                   ("Workover", "Kill-fluid loss"),
                   ("Production", "Scale, fines,\nasphaltenes")],
               y=2.18, h=2.35, with_rail=True, rail_y=1.98)
    for i in range(5):
        w = (CONTENT_W - 4 * 0.28) / 5
        x = M + i * (w + 0.28) + w / 2
        text(s, str(i + 1), x - 0.135, 1.85, 0.27, 0.3, size=11, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER)
    rect(s, M, 4.95, CONTENT_W, 0.78, fill=CREAM, line=AMBER, lw=1.5)
    text(s, "This talk focuses on drilling-fluid damage", M, 5.16, CONTENT_W, 0.4,
         size=18, color=AMBERD, bold=True, align=PP_ALIGN.CENTER)
    source(s, "Adapted from Krueger (1986).")
    pagenum(s, 8)
    notes(s, "Damage is not only a drilling problem: every stage from drilling to production "
             "can impair the near-wellbore zone. This talk focuses on drilling-fluid damage, "
             "the earliest and most preventable source.")
    return s


# ============================================================ 9  INVASION
def s09(prs):
    s = blank(prs)
    heading(s, "Drilling-fluid invasion is the primary source of damage")
    tracker(s, 2)
    bullets_from(s, ["Overbalance pushes fluid into the rock",
                     "Spurt loss, then filtrate, then filter cake",
                     "Severity grows with overbalance and exposure time",
                     "Mud properties and permeability also matter"],
                M, 2.00, 4.80, 4.40, size=17)
    pic(s, "s9_invasion.png", 5.60, 1.98, 6.85)
    source(s, "Schematic by author.")
    pagenum(s, 9)
    notes(s, "Invasion is driven by overbalance: mud pressure above formation pressure. "
             "The sequence is spurt loss, then filtrate invasion, while a filter cake builds. "
             "Severity grows with overbalance and exposure time, and depends on mud "
             "properties and formation permeability.")
    return s


# ============================================================ 10  SOLIDS AND FINES
def s10(prs):
    s = blank(prs)
    heading(s, "Solids and fines plug pore throats")
    tracker(s, 2)
    bullets_from(s, ["Mud solids bridge or enter pore throats",
                     "Formation fines detach and migrate",
                     "Result: severe, hard-to-reverse permeability loss"],
                M, 2.00, 4.80, 4.40, size=17)
    pic(s, "s10_solids.png", 5.60, 1.98, 6.85)
    source(s, "Schematic by author. [Or: SEM image, source, year]")
    pagenum(s, 10)
    notes(s, "Two solid-related mechanisms: mud solids bridge at the pore throats or invade "
             "deeper, and formation fines detach and migrate until they block a throat. "
             "Both can cause severe permeability loss that is hard to reverse.")
    return s


# ============================================================ 11  CLAY
def s11(prs):
    s = blank(prs)
    heading(s, "Incompatible water makes clays swell and migrate")
    tracker(s, 2)
    data = [["Clay", "Main damage behaviour"],
            ["Smectite", "Swells, blocks pores"],
            ["Illite", "Migrates as fines"],
            ["Kaolinite", "Detaches and migrates"],
            ["Chlorite", "Iron precipitates with acid"]]
    add_table(s, data, M, 2.00, 6.95, [2.40, 4.55], row_h=0.72, header_h=0.52,
              size=14, header_size=14)
    bullets_from(s, ["Fresh filtrate triggers swelling and dispersion",
                     "Salinity and ion composition matter",
                     "Fix: KCl brines, clay stabilizers"],
                7.85, 2.10, 4.90, 3.0, size=16)
    source(s, "Simplified from Civan, Reservoir Formation Damage.")
    pagenum(s, 11)
    notes(s, "Clay chemistry drives water sensitivity. Smectite swells, illite and kaolinite "
             "migrate as fines, chlorite releases iron when acidised. The fixes are "
             "inhibitive brines such as KCl and clay stabilisers.")
    return s


# ============================================================ 12  CAPILLARY
def s12(prs):
    s = blank(prs)
    heading(s, "Capillary forces trap filtrate near the wellbore")
    tracker(s, 2)
    bullets_from(s, ["Water block cuts hydrocarbon relative permeability",
                     "Strongest in low-permeability rock",
                     "Emulsions and wettability change add resistance"],
                M, 2.05, 4.65, 2.6, size=16)
    rect(s, M, 4.25, 4.65, 2.15, fill=PANEL2)
    text(s, "P~c~ = 2σ cosθ / r", M, 4.55, 4.65, 0.5, size=25, color=NAVY, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, "(σ = interfacial tension, θ = contact angle, r = pore-throat radius)",
         M + 0.25, 5.25, 4.15, 0.8, size=12, color=GREY, align=PP_ALIGN.CENTER, line=1.25)
    pic(s, "s12_capillary.png", 5.55, 2.70, 7.25)
    source(s, "Schematic by author.")
    pagenum(s, 12)
    notes(s, "Capillary forces trap filtrate in the pore space. A water block cuts the "
             "relative permeability to hydrocarbon, worst in low-permeability rock, where "
             "small pore throats give high capillary pressure — hence the equation.")
    return s


# ============================================================ 13  PRECIPITATES
def s13(prs):
    s = blank(prs)
    heading(s, "Fluid incompatibility creates precipitates in the pores")
    tracker(s, 2)
    bullets_from(s, ["Filtrate + brine can form CaCO~3~, CaSO~4~, BaSO~4~",
                     "Acids or high-pH fluids: gels, sludges",
                     "Prevent: lab compatibility tests before pumping"],
                M, 2.00, 4.80, 4.40, size=17)
    pic(s, "s13_precip.png", 5.45, 2.70, 7.40)
    source(s, "Schematic by author.")
    pagenum(s, 13)
    notes(s, "Incompatible fluids can precipitate solids in the pore space: calcium carbonate, "
             "calcium sulphate, barium sulphate, and gels or sludges with acids or high-pH "
             "fluids. Run compatibility tests before pumping.")
    return s


# ============================================================ 14  DETECTION
def s14(prs):
    s = blank(prs)
    heading(s, "Damage is detected by tests and production data")
    tracker(s, 3)
    data = [["Method", "What it shows"],
            ["Pressure buildup test", "Total skin"],
            ["PI versus offset wells", "Underperformance"],
            ["Core flow test", "Return permeability"],
            ["Production logging", "Inflow profile"],
            ["Nodal analysis", "Expected versus actual rate"]]
    add_table(s, data, M, 2.00, 6.45, [2.75, 3.70], row_h=0.60, header_h=0.50,
              size=13.5, header_size=13.5)
    rect(s, M, 5.45, 6.45, 1.10, fill=PANEL, line=BORDER, lw=1.0)
    box, tf = textbox(s, M + 0.22, 5.62, 6.05, 0.85)
    rich(para(tf, first=True, line=1.25, after=2),
         "Return permeability = k~after~ / k~before~", 14, NAVY, True)
    rich(para(tf, line=1.25), "Test skin is total skin, not only damage", 13, GREY)
    pic(s, "s14_buildup.png", 7.05, 2.20, 5.70)
    source(s, "Concept sketch by author. After Earlougher (1977).")
    pagenum(s, 14)
    notes(s, "Detection toolkit: pressure-buildup analysis gives total skin, offset-well "
             "comparison and production logging locate the loss, core flow tests give return "
             "permeability, and nodal analysis compares expected with actual rate. Note that "
             "test skin is total skin.")
    return s


# ============================================================ 15  HAWKINS
def s15(prs):
    s = blank(prs)
    heading(s, "Hawkins' equation links the damaged zone to skin")
    tracker(s, 3)
    rect(s, M, 1.95, CONTENT_W, 1.15, fill=PANEL2, line=BORDER, lw=1.0)
    text(s, "[red]**S**[/red]  =  (k/k~s~ − 1) ln(r~s~/r~w~)", M, 2.28, CONTENT_W, 0.6,
         size=27, color=NAVY, align=PP_ALIGN.CENTER)
    pic(s, "s15_hawkins.png", M + 0.05, 3.45, 3.85)
    legend = [("k", "undamaged permeability (md)"),
              ("k~s~", "damaged permeability (md)"),
              ("r~s~", "damage radius (ft)"),
              ("r~w~", "wellbore radius (ft)")]
    box, tf = textbox(s, 4.60, 3.45, 3.10, 3.20)
    for i, (sym, desc) in enumerate(legend):
        rich(para(tf, first=(i == 0), line=1.15, after=3), sym, 16, NAVY, True)
        rich(para(tf, line=1.2, after=11), desc, 12.5, GREY)
    bullets_from(s, ["Severity (k/k~s~) acts linearly",
                     "Depth (r~s~) acts logarithmically",
                     "Assumes a uniform, cylindrical damaged zone",
                     "Well-test skin is total skin"],
                7.95, 3.45, 4.80, 3.2, size=16)
    source(s, "Hawkins (1956), Trans. AIME 207.", y=6.95)
    pagenum(s, 15)
    notes(s, "Hawkins' equation is the workhorse: the severity of the permeability reduction "
             "acts linearly and the depth of the damaged zone enters logarithmically. Two "
             "cautions: the model assumes a uniform cylindrical damaged zone, and well-test "
             "skin is total skin, not only damage.")
    return s


# ============================================================ 16  WORKED EXAMPLE
def s16(prs):
    s = blank(prs)
    heading(s, "Worked example: a skin of 8.5 costs 55% of productivity", p=30, w=9.35)
    tag(s)
    data = [["INPUT", "OUTPUT"],
            ["k, k~s~: 100, 20 md", "S = 8.55"],
            ["h = 50 ft", "q (S = 0) = 4,270 STB/d"],
            ["r~w~, r~e~: 0.354, 745 ft", "q (damaged) = 1,910 STB/d"],
            ["r~s~ = 3 ft", "FE = 0.447"],
            ["p̄, p~wf~: 3,500, 2,500 psi", "Δp~skin~ = 553 psi"],
            ["μ = 1 cP, B = 1.2 rb/STB", ""]]
    add_table(s, data, M, 2.00, 6.35, [3.35, 3.00], row_h=0.62, header_h=0.50,
              size=13.5, header_size=13.5)
    pic(s, "s16_bars.png", 7.00, 2.40, 5.75)
    source(s, "Author's calculation using Hawkins (1956) and radial Darcy flow.")
    pagenum(s, 16)
    notes(s, "Worked example with illustrative data: a permeability ratio of five over a 3 ft "
             "zone gives a skin of 8.55, which cuts the rate from 4,270 to 1,910 STB/d — "
             "55% of productivity lost. The skin pressure drop is 553 psi of the 1,000 psi "
             "drawdown.")
    return s


# ============================================================ 17  SEVERITY VS DEPTH
def s17(prs):
    s = blank(prs)
    heading(s, "Severity of permeability reduction matters more than depth", p=30, w=9.35)
    tag(s)
    data = [["k/k~s~", "r~s~ = 1 ft", "r~s~ = 3 ft", "r~s~ = 5 ft"],
            ["2", "1.0 / 0.87", "2.1 / 0.76", "2.6 / 0.72"],
            ["5", "4.2 / 0.62", "8.5 / 0.45", "10.6 / 0.39"],
            ["10", "9.3 / 0.42", "19.2 / 0.26", "23.8 / 0.22"]]
    add_table(s, data, M, 2.00, 6.45, [1.65, 1.60, 1.60, 1.60], row_h=0.62,
              header_h=0.55, size=14, header_size=14,
              align_center_cols=(1, 2, 3))
    rect(s, M, 4.50, 6.45, 1.80, fill=CREAM, line=AMBER, lw=1.5)
    box, tf = textbox(s, M + 0.25, 4.67, 5.95, 1.45)
    rich(para(tf, first=True, line=1.25, after=7), "Depth 3 → 5 ft:   FE 0.45 → 0.39",
         14.5, NAVY, True)
    rich(para(tf, line=1.25), "Severity 2 → 10:   FE 0.76 → 0.26", 14.5, NAVY, True)
    text(s, "S = skin  ·  FE = flow efficiency", M + 0.25, 5.72, 5.95, 0.4, size=11.5,
         color=GREY)
    pic(s, "s17_fe.png", 6.95, 2.15, 5.80)
    source(s, "Author's calculation. Same reservoir as slide 16.", y=6.55)
    pagenum(s, 17)
    notes(s, "Sensitivity study: growing the damage radius from 3 to 5 ft moves flow "
             "efficiency from 0.45 to 0.39, while worsening severity from 2 to 10 moves it "
             "from 0.76 to 0.26. Severity dominates — so aim to recover permeability rather "
             "than to chase the last foot of penetration.")
    return s


# ============================================================ 18  PREVENTION
def s18(prs):
    s = blank(prs)
    heading(s, "Prevention starts with fluid design")
    tracker(s, 4)
    bullets_from(s, ["Size bridging agents ≈ ⅓ of pore throats",
                     "Use acid-soluble CaCO~3~ bridging particles",
                     "Low filtrate loss: thin, tough filter cake",
                     "Inhibitive brines (KCl) protect clays",
                     "Minimize overbalance and exposure time"],
                M, 2.00, 4.85, 4.40, size=17)
    pic(s, "s18_fluids.png", 5.60, 2.05, 6.95)
    source(s, "Schematic by author. Bridging rule after Abrams (1977).")
    pagenum(s, 18)
    notes(s, "Prevention is really fluid design: bridging particles sized at about one third "
             "of the pore throats, acid-soluble carbonate bridging solids, low filtrate loss, "
             "inhibitive brines such as KCl, and minimum overbalance and exposure time.")
    return s


# ============================================================ 19  REMEDIATION
def s19(prs):
    s = blank(prs)
    heading(s, "Remediation must match the damage type")
    tracker(s, 4)
    data = [["Damage", "Typical treatment"],
            ["Carbonate solids, scale", "HCl"],
            ["Clay and silicate fines", "Mud acid (HF/HCl)"],
            ["Emulsion, organics, paraffin", "Solvents, surfactants"],
            ["Deep or severe damage", "Hydraulic fracturing"]]
    add_table(s, data, M, 2.00, 6.20, [3.05, 3.15], row_h=0.70, header_h=0.50,
              size=13.5, header_size=13.5)
    rect(s, M, 5.45, 6.20, 0.85, fill=CREAM, line=AMBER, lw=1.5)
    text(s, "Diagnose first. Wrong treatment creates new damage.", M + 0.2, 5.67,
         5.85, 0.45, size=14.5, color=AMBERD, bold=True, align=PP_ALIGN.CENTER)
    pic(s, "s19_tree.png", 6.95, 2.15, 5.80)
    source(s, "Simplified from Economides et al. Always verify with lab tests.")
    pagenum(s, 19)
    notes(s, "Remediation must match the damage: acid for carbonate solids and scale, mud "
             "acid for clays and fines, solvents for organics, fracturing for deep or severe "
             "damage. Diagnose first — the wrong chemistry creates new damage.")
    return s


# ============================================================ 20  ECONOMICS
def s20(prs):
    s = blank(prs)
    heading(s, "Remediation can pay back fast, but prevention avoids cost and risk", p=30, w=9.35)
    tag(s)
    bullets_from(s, ["Acidizing example: S falls from 8.55 to 2",
                     "Rate rises 1,910 → 3,310 STB/d",
                     "Payback ≈ 14 days (assumed values)",
                     "HSE: acid handling, flowback disposal",
                     "Prevention avoids both cost and risk"],
                M, 2.00, 5.60, 2.95, size=16.5)
    rect(s, M, 5.15, 5.60, 1.25, fill=PANEL, line=BORDER, lw=1.0)
    box, tf = textbox(s, M + 0.24, 5.40, 5.15, 0.9)
    rich(para(tf, first=True, line=1.35, after=3),
         "**Assumed:** margin $30/bbl · 50% of gain sustained", 12.5, GREY)
    rich(para(tf, line=1.35),
         "job cost $300,000 · net gain ≈ $21,075/day", 12.5, GREY)
    pic(s, "s20_econ.png", 6.75, 2.35, 6.00)
    source(s, "Author's calculation. Economic inputs are assumptions.")
    pagenum(s, 20)
    notes(s, "Economics with assumed inputs: acidising takes skin from 8.55 to 2, adding "
             "1,405 STB/d; at a $30 per barrel margin with half of the gain sustained, the "
             "job pays back in about 14 days. Add the HSE side: acid handling and flowback "
             "disposal.")
    return s


# ============================================================ 21  PROJECT
def s21(prs):
    s = blank(prs)
    heading(s, "Next step: a skin-based screening model for drilling-fluid damage", p=30)
    tracker(s, 5)
    five_boxes(s, [("Inputs", "k, h, k~s~, r~s~, drawdown"),
                   ("Skin", "Hawkins equation"),
                   ("Productivity", "q and FE"),
                   ("Economics", "Payback, net value"),
                   ("Decision", "Prevent, treat, or ignore")],
               y=2.30, h=2.15, with_rail=True, rail_y=2.00)
    for i in range(5):
        w = (CONTENT_W - 4 * 0.28) / 5
        x = M + i * (w + 0.28) + w / 2
        text(s, str(i + 1), x - 0.135, 1.87, 0.27, 0.3, size=11, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER)
    rect(s, M, 4.95, CONTENT_W, 1.70, fill=PANEL, line=BORDER, lw=1.0)
    box, tf = textbox(s, M + 0.30, 5.15, CONTENT_W - 0.6, 1.4)
    rows = [("Gap", "damage is often treated qualitatively"),
            ("Tools", "Excel or Python, literature return-permeability data"),
            ("Output", "sensitivity charts and a decision guide")]
    for i, (lab, txt) in enumerate(rows):
        rich(para(tf, first=(i == 0), line=1.2, after=6), f"**{lab}:**  {txt}", 14.5, INK)
    pagenum(s, 21)
    notes(s, "The graduation project: a spreadsheet or Python screening model that takes k, h, "
             "ks, rs and drawdown, applies Hawkins, predicts rate and flow efficiency, adds "
             "simple economics, and outputs a prevent / treat / ignore recommendation.")
    return s


# ============================================================ 22  CONCLUSIONS
def s22(prs):
    s = blank(prs)
    heading(s, "Conclusions: damage is quantifiable and largely preventable", p=32)
    items = [("1. Mechanisms", ["Invasion damages by plugging, swelling, trapping"]),
             ("2. Quantification", ["Example: skin 8.5 costs 55%",
                                    "Severity outweighs depth"]),
             ("3. Strategies", ["Prevention beats remediation",
                                "Match treatment to damage type"])]
    gap = 0.235
    w = (CONTENT_W - 2 * gap) / 3
    for i, (head, bl) in enumerate(items):
        x = M + i * (w + gap)
        rect(s, x, 2.20, w, 3.90, fill=WHITE, line=BORDER, lw=1.25)
        rect(s, x, 2.20, w, 0.09, fill=TEAL)
        text(s, head, x + 0.25, 2.55, w - 0.5, 0.45, size=19, color=NAVY, bold=True)
        bullets_from(s, bl, x + 0.25, 3.20, w - 0.5, 2.55, size=16,
                     anchor=MSO_ANCHOR.TOP)
    pagenum(s, 22)
    notes(s, "Three conclusions. First: invasion damages by plugging, swelling and trapping. "
             "Second: a skin of 8.5 costs 55% of rate, and severity outweighs depth. Third: "
             "prevention beats remediation, and any treatment must match the damage type.")
    return s


# ============================================================ 23  REFERENCES
def s23(prs):
    s = blank(prs)
    heading(s, "References and questions")
    refs = ["Hawkins, M.F. (1956). A note on the skin effect. Trans. AIME, 207.",
            "van Everdingen, A.F. (1953). The skin effect and its influence on the "
            "productive capacity of a well. Trans. AIME, 198.",
            "Krueger, R.F. (1986). An overview of formation damage and well productivity "
            "in oilfield operations. JPT, 38(2).",
            "Abrams, A. (1977). Mud design to minimize rock impairment due to particle "
            "invasion. JPT, 29(5).",
            "Earlougher, R.C. (1977). Advances in Well Test Analysis. SPE Monograph 5.",
            "Civan, F. Reservoir Formation Damage. Gulf Professional Publishing.",
            "Economides, M.J. et al. Petroleum Production Systems. Prentice Hall.",
            "Bourgoyne, A.T. et al. (1986). Applied Drilling Engineering. SPE.",
            "[Recent papers, 2018–2025: Author (Year), Title, Journal — items 9–12]"]
    box, tf = textbox(s, M, 1.58, 7.90, 4.95)
    for i, r in enumerate(refs):
        p = para(tf, first=(i == 0), line=1.18, after=7,
                 indent=0.26 if i < 8 else 0.0)
        rich(p, f"{i+1}.\t{r}" if i < 8 else r, 12, INK if i < 8 else AMBERD,
             bold=(i >= 8))
    line(s, 8.85, 1.58, 8.85, 6.45, BORDER, 1.25)
    text(s, "Thank you.", 9.25, 2.55, 3.60, 0.6, size=34, color=NAVY, bold=True)
    text(s, "Questions?", 9.25, 3.30, 3.60, 0.6, size=34, color=TEALD, bold=True)
    text(s, "[your.email@university]", 9.25, 4.30, 3.60, 0.4, size=14, color=GREY)
    text(s, "Verify every reference in your library (volume, pages, editions) before "
            "submitting.", M, 6.55, 7.90, 0.6, size=11, color=GREY, italic=True, line=1.25)
    pagenum(s, 23)
    notes(s, "Thank the audience, point to the reference list, and invite questions. Note "
             "that items 9 to 12 must come from your own literature table before submitting.")
    return s


# ============================================================ BACKUPS
def b1(prs):
    s = blank(prs)
    heading(s, "Backup: Hawkins' formula comes from series radial resistances", p=27)
    rect(s, M, 1.85, CONTENT_W, 4.10, fill=PANEL2, line=BORDER, lw=1.0)
    box, tf = textbox(s, M + 0.5, 2.15, CONTENT_W - 1.0, 3.6)
    rows = [("Damaged zone:", "Δp~s~ = [ q B μ / (0.00708 k~s~ h) ] ln(r~s~/r~w~)"),
            ("Undamaged zone:", "Δp~u~ = [ q B μ / (0.00708 k h) ] [ ln(r~e~/r~s~) − 0.75 ]"),
            ("Sum:", "ln(r~e~/r~w~) − 0.75 + (k/k~s~ − 1) ln(r~s~/r~w~)"),
            ("Skin term:", "[red]**S**[/red] = (k/k~s~ − 1) ln(r~s~/r~w~)")]
    for i, (lab, eq) in enumerate(rows):
        p = para(tf, first=(i == 0), line=1.2, after=22)
        rich(p, f"**{lab}**   ", 17, NAVY)
        rich(p, eq, 20, INK)
    text(s, "Both zones are in series, so the pressure drops add — and the extra term "
            "is the skin.", M + 0.5, 6.15, CONTENT_W - 1.0, 0.5, size=13, color=GREY,
         italic=True)
    text(s, "B1", SLIDE_W - M - 0.5, SLIDE_H - 0.52, 0.5, 0.3, size=11, color=LGREY,
         bold=True, align=PP_ALIGN.RIGHT)
    notes(s, "Backup: the derivation of Hawkins' equation from two radial resistances in "
             "series — the damaged zone and the undamaged zone.")
    return s


def b2(prs):
    s = blank(prs)
    heading(s, "Backup: test skin combines several effects", p=27)
    rect(s, M, 2.10, CONTENT_W, 1.45, fill=PANEL2, line=BORDER, lw=1.0)
    text(s, "S~total~ = S~damage~ + S~perforation~ + S~partial penetration~ "
            "+ S~deviation~ + S~non-Darcy~", M + 0.3, 2.62, CONTENT_W - 0.6, 0.6,
         size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    rect(s, M, 3.95, CONTENT_W, 0.85, fill=CREAM, line=AMBER, lw=1.5)
    text(s, "Only S~damage~ is removable by damage treatment", M, 4.15, CONTENT_W, 0.45,
         size=18, color=AMBERD, bold=True, align=PP_ALIGN.CENTER)
    text(s, "A high test skin is not automatically a formation-damage problem — check "
            "completion and geometry first.", M, 5.15, CONTENT_W, 0.5, size=14, color=GREY,
         italic=True, align=PP_ALIGN.CENTER)
    text(s, "B2", SLIDE_W - M - 0.5, SLIDE_H - 0.52, 0.5, 0.3, size=11, color=LGREY,
         bold=True, align=PP_ALIGN.RIGHT)
    notes(s, "Backup: total test skin is the sum of damage, perforation, partial penetration, "
             "deviation and non-Darcy effects. Only the damage part can be removed by a "
             "damage treatment.")
    return s


def b3(prs):
    s = blank(prs)
    heading(s, "Backup: bridging theory sets particle size", p=27)
    bullets_from(s, ["Median bridging particle ≈ ⅓ median pore throat",
                     "Bridging solids ≥ 5% by volume (Abrams)",
                     "Too fine: deep invasion. Too coarse: no bridge"],
                M, 2.20, 11.0, 2.6, size=20)
    rect(s, M, 4.60, CONTENT_W, 1.20, fill=PANEL, line=BORDER, lw=1.0)
    text(s, "Sizing is a balance: fine enough to form a bridge, coarse enough to stay "
            "on the wall.", M + 0.3, 4.95, CONTENT_W - 0.6, 0.6, size=14, color=GREY,
         italic=True)
    text(s, "B3", SLIDE_W - M - 0.5, SLIDE_H - 0.52, 0.5, 0.3, size=11, color=LGREY,
         bold=True, align=PP_ALIGN.RIGHT)
    notes(s, "Backup: the Abrams bridging rule — bridging particles at about one third of "
             "the median pore-throat size, with enough solids to build the bridge.")
    return s


def b4(prs):
    s = blank(prs)
    heading(s, "Backup: drawdown scales rate, not flow efficiency", p=27)
    bullets_from(s, ["q scales linearly with drawdown",
                     "FE depends only on ln(r~e~/r~w~) − 0.75 and S",
                     "Relative loss is independent of drawdown"],
                M, 2.10, 5.60, 2.8, size=17)
    pic(s, "b4_drawdown.png", 6.85, 2.05, 5.90)
    text(s, "B4", SLIDE_W - M - 0.5, SLIDE_H - 0.52, 0.5, 0.3, size=11, color=LGREY,
         bold=True, align=PP_ALIGN.RIGHT)
    notes(s, "Backup: increasing drawdown lifts both rates proportionally, so flow efficiency "
             "stays at 0.447 — you cannot produce your way out of damage.")
    return s


def b5(prs):
    s = blank(prs)
    heading(s, "Backup: fluid comparison (qualitative)", p=27)
    data = [["Fluid", "Main advantage", "Main damage risk"],
            ["Water-based", "Low cost", "Clay swelling, water block"],
            ["Oil-based", "Shale inhibition", "Wettability change, emulsions"],
            ["Formate brine", "Low solids", "High cost"],
            ["Polymer + CaCO~3~", "Thin, acid-soluble cake", "Polymer residue"]]
    add_table(s, data, M, 2.15, CONTENT_W, [3.40, 4.20, 4.57], row_h=0.80,
              header_h=0.55, size=14.5, header_size=14.5)
    text(s, "Check this table against your course notes. The risks depend on the formation.",
         M, 5.85, CONTENT_W, 0.5, size=13, color=GREY, italic=True)
    text(s, "B5", SLIDE_W - M - 0.5, SLIDE_H - 0.52, 0.5, 0.3, size=11, color=LGREY,
         bold=True, align=PP_ALIGN.RIGHT)
    notes(s, "Backup: a qualitative comparison of drilling-fluid families, with the main "
             "advantage and the main damage risk of each.")
    return s


# ============================================================ MAIN
def main():
    prs = prs_new()
    builders = [s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12, s13,
                s14, s15, s16, s17, s18, s19, s20, s21, s22, s23,
                b1, b2, b3, b4, b5]
    slides = [f(prs) for f in builders]
    for s in slides[23:]:
        hide(prs, s)
    prs.core_properties.title = "Formation Damage During Drilling"
    prs.core_properties.author = "[Your Name]"
    prs.core_properties.subject = ("Mechanisms, skin quantification and mitigation "
                                   "strategies — seminar deck")
    out = "Formation_Damage_Drilling.pptx"
    prs.save(out)
    print(f"saved {out}: {len(list(prs.slides))} slides "
          f"({len(list(prs.slides)) - 5} visible + 5 hidden backups)")


if __name__ == "__main__":
    main()
