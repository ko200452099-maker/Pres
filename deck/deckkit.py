"""Helpers for building the Formation Damage deck: geometry, text, rich runs, tables."""
import glob
import math
import re
import copy
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ----------------------------------------------------------------- palette
NAVY  = RGBColor(0x0B, 0x2A, 0x4A)
TEAL  = RGBColor(0x12, 0xA4, 0xA8)
TEALD = RGBColor(0x0E, 0x7C, 0x7C)
RED   = RGBColor(0xD1, 0x48, 0x3B)
AMBER = RGBColor(0xE3, 0x9A, 0x26)
AMBERD= RGBColor(0x8C, 0x6D, 0x10)
GREY  = RGBColor(0x5A, 0x64, 0x72)
LGREY = RGBColor(0xB9, 0xC2, 0xCC)
BORDER= RGBColor(0xDC, 0xE3, 0xEA)
PANEL = RGBColor(0xF2, 0xF5, 0xF9)
PANEL2= RGBColor(0xEA, 0xF1, 0xF6)
CREAM = RGBColor(0xFD, 0xF3, 0xE0)
TEALP = RGBColor(0xE6, 0xF5, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK   = RGBColor(0x1F, 0x2A, 0x37)

COLORS = {"navy": NAVY, "teal": TEAL, "teald": TEALD, "red": RED, "amber": AMBER,
          "amberd": AMBERD, "grey": GREY, "lgrey": LGREY, "ink": INK,
          "white": WHITE, "panel": PANEL}

FONT = "Calibri"

SLIDE_W, SLIDE_H = 13.333, 7.5
M = 0.58                     # page margin
CONTENT_W = SLIDE_W - 2 * M  # 12.173

# ------------------------------------------------- measurement (Calibri metrics via Carlito)
_FONT_CACHE = {}
_CARLITO = (glob.glob("/usr/share/fonts/**/Carlito-Regular.ttf", recursive=True) or
            glob.glob("/usr/share/fonts/**/DejaVuSans.ttf", recursive=True))[0]


def _font(size_pt):
    from PIL import ImageFont
    key = round(size_pt * 4)
    if key not in _FONT_CACHE:
        _FONT_CACHE[key] = ImageFont.truetype(_CARLITO, key)
    return _FONT_CACHE[key]


def text_w(s, size_pt):
    """Width of string in points."""
    return _font(size_pt).getlength(s) / 4.0


def strip_markup(s):
    return re.sub(r"(\*\*|~|\^|\[/?[a-z]+\])", "", s)


def wrapped_lines(s, width_in, size_pt):
    """Number of wrapped lines for text in a box of given width (inches)."""
    s = strip_markup(s)
    limit = width_in * 72.0
    words = s.split()
    lines, cur = 1, ""
    for w in words:
        trial = w if not cur else cur + " " + w
        if text_w(trial, size_pt) <= limit or not cur:
            cur = trial
        else:
            lines += 1
            cur = w
    return lines


def block_height(items, width_in, size_pt, line=1.18, space=8.0, indent=0.0):
    """Estimated height (inches) of a bullet block."""
    h = 0.0
    for it in items:
        n = wrapped_lines(it, width_in - indent, size_pt)
        h += (n * size_pt * line + space) / 72.0
    return h


def fit_size(items, width_in, height_in, p=16.0, minimum=11.0, line=1.18,
             space=8.0, indent=0.0):
    """Largest font size <= p at which the bullet block fits the given height."""
    size = p
    while size > minimum:
        if block_height(items, width_in, size, line, space, indent) <= height_in:
            return size
        size -= 0.5
    return minimum


# ----------------------------------------------------------------- primitives
def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
         adj=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if adj is not None:
        sh.adjustments[0] = adj
    sh.text_frame.word_wrap = True
    return sh


def line(slide, x1, y1, x2, y2, color=LGREY, lw=1.0, dash=None):
    from pptx.enum.shapes import MSO_CONNECTOR
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(lw)
    if dash:
        cn.line.dash_style = dash
    return cn


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return box, tf


def para(tf, first=False, align=PP_ALIGN.LEFT, line=1.2, before=0, after=0,
         indent=0.0, bullet=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if indent:
        p._pPr.set("marL", str(int(indent * 914400)))
        p._pPr.set("indent", str(int(-indent * 914400)))
    return p


def rich(p, markup, size, color=INK, bold=False, italic=False, font=FONT):
    """Add runs to paragraph p, honouring **bold**, ~sub~, ^sup^, [colour]...[/colour]."""
    fmt = {"bold": bold, "color": color, "sub": False, "sup": False}
    stack = []
    for t in re.split(r"(\*\*|~|\^|\[/?[a-z]+\])", markup):
        if not t:
            continue
        if t == "**":
            fmt["bold"] = not fmt["bold"]
        elif t == "~":
            fmt["sub"] = not fmt["sub"]
        elif t == "^":
            fmt["sup"] = not fmt["sup"]
        elif t.startswith("[/"):
            if stack:
                fmt.update(stack.pop())
        elif t.startswith("[") and t.endswith("]"):
            name = t[1:-1]
            if name in COLORS:
                stack.append(copy.copy(fmt))
                fmt["color"] = COLORS[name]
        else:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(size)
            f.bold = fmt["bold"]
            f.italic = italic
            f.name = font
            f.color.rgb = fmt["color"]
            if fmt["sub"]:
                r.font._rPr.set("baseline", "-25000")
            elif fmt["sup"]:
                r.font._rPr.set("baseline", "30000")
    return p


def text(slide, markup, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         line=1.2, anchor=MSO_ANCHOR.TOP, italic=False, after=0, before=0):
    box, tf = textbox(slide, x, y, w, h, anchor)
    rich(para(tf, first=True, align=align, line=line, after=after, before=before),
         markup, size, color, bold, italic)
    return box


def bullet_lines(slide, items, x, y, w, h, size=16, color=INK, line=1.18, space=8,
                 bullet_color=None, bullet_char="•", align=PP_ALIGN.LEFT,
                 indent=0.20, bold_first=False):
    """Bulleted block. Size is reduced automatically until the block fits the box."""
    indent = 0.20
    size = fit_size(items, w, h, p=size, line=line, space=space, indent=indent)
    box, tf = textbox(slide, x, y, w, h)
    for i, it in enumerate(items):
        p = para(tf, first=(i == 0), line=line, after=space, indent=indent,
                 align=align)
        # bullet glyph
        r = p.add_run()
        r.text = bullet_char + "\t"
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = bullet_color or TEAL
        r.font.bold = True
        rich(p, it, size, color, bold=(bold_first and i == 0))
        # tab stop so text aligns after the glyph
        p._pPr.set("marL", str(int(indent * 914400)))
        p._pPr.set("indent", str(int(-indent * 914400)))
    return box, size


def bullets_from(slide, items, x, y, w, h, size=16, space=9,
                 anchor=MSO_ANCHOR.MIDDLE):
    """Bullet block, vertically centred in its box, auto-sized to fit."""
    size = fit_size(items, w - 0.2, h, p=size, space=space, indent=0.2)
    box, tf = textbox(slide, x, y, w, h, anchor)
    for i, it in enumerate(items):
        p = para(tf, first=(i == 0), line=1.18, after=space, indent=0.22)
        r = p.add_run()
        r.text = "•\t"
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = TEAL
        r.font.bold = True
        rich(p, it, size, INK)
    return size


# ----------------------------------------------------------------- tables
def add_table(slide, data, x, y, w, col_w, row_h=None, header_h=0.5,
              size=13.5, header_size=None, header_fill=NAVY, zebra=True,
              first_col_bold=False, align_center_cols=()):
    rows, cols = len(data), len(data[0])
    total_h = header_h + (rows - 1) * (row_h or 0.45)
    shp = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w),
                                 Inches(total_h))
    tbl = shp.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    tbl.rows[0].height = Inches(header_h)
    for i in range(1, rows):
        tbl.rows[i].height = Inches(row_h or 0.45)
    hs = header_size or size
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.10)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.045)
            cell.margin_bottom = Inches(0.045)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
            elif zebra and i % 2 == 0:
                cell.fill.fore_color.rgb = PANEL
            else:
                cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.05
            if j in align_center_cols and i > 0:
                p.alignment = PP_ALIGN.CENTER
            col = WHITE if i == 0 else (NAVY if (j == 0 or (first_col_bold and i > 0)) else INK)
            rich(p, val, hs if i == 0 else size, col,
                 bold=(i == 0 or j == 0 or first_col_bold))
    return tbl


# ----------------------------------------------------------------- deck furniture
def title(slide, text_markup, size=30, y=0.34, h=0.9, w=None, color=NAVY):
    w = w or CONTENT_W
    box, tf = textbox(slide, M, y, w, h)
    rich(para(tf, first=True, line=1.02), text_markup, size, color, bold=True)
    return box


def source(slide, markup, y=None):
    y = y if y is not None else SLIDE_H - 0.52
    text(slide, markup, M, y, CONTENT_W - 0.9, 0.30, size=10.5, color=GREY, italic=True)


def pagenum(slide, n):
    text(slide, str(n), SLIDE_W - 0.95, SLIDE_H - 0.52, 0.5, 0.30, size=11,
         color=LGREY, align=PP_ALIGN.RIGHT, bold=True)


SECTIONS = [
    ("Fundamentals", "Flow near the well", 1),
    ("Mechanisms", "How damage forms", 2),
    ("Quantification", "Skin and productivity", 3),
    ("Solutions", "Prevention and cure", 4),
    ("Next steps", "Economics and project", 5),
]


def tracker(slide, seg):
    """Progress tracker: five short segments, active one filled."""
    x0, y = M, 1.30
    gap = 0.28
    w = (CONTENT_W - 4 * gap) / 5
    for i, (name, _, idx) in enumerate(SECTIONS, start=1):
        x = x0 + (i - 1) * (w + gap)
        active = (i == seg)
        rect(slide, x, y, w, 0.075, fill=TEAL if active else RGBColor(0xDF, 0xE6, 0xEC))
        text(slide, name, x, y + 0.13, w, 0.26, size=10.5,
             color=TEAL if active else LGREY, bold=active)


def tag(slide, label="ILLUSTRATIVE DATA", color=AMBERD, fill=CREAM, w=2.15):
    rect(slide, SLIDE_W - M - w, 0.40, w, 0.36, fill=fill, line=AMBER, lw=1.0)
    text(slide, label, SLIDE_W - M - w, 0.445, w, 0.28, size=11, color=color,
         bold=True, align=PP_ALIGN.CENTER)


def notes(slide, text_str):
    slide.notes_slide.notes_text_frame.text = text_str


def hide(prs, slide):
    """Mark a slide as hidden in PowerPoint (it stays in the file)."""
    idx = list(prs.slides).index(slide)
    sldId = list(prs.slides._sldIdLst)[idx]
    sldId.set("show", "0")
